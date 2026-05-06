"""
Styler — única fonte de estilos para os arquivos Excel.

Todas as funções retornam objetos openpyxl prontos para aplicar a células.
O writer.py usa este módulo; nenhuma outra camada aplica estilos diretamente.
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from utils.constants import (
    COL_HEADER_BG, COL_HEADER_FG,
    ROW_URGENCY_BG, SECTION_BG, SECTION_FG,
    STATUS_CELL_BG, URGENCY_CELL_BG,
)


# ── Preenchimentos ────────────────────────────────────────────────────────────

def fill(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=hex_color)


def section_fill() -> PatternFill:
    return fill(SECTION_BG)


def col_header_fill() -> PatternFill:
    return fill(COL_HEADER_BG)


def row_fill(urgency: str) -> PatternFill:
    return fill(ROW_URGENCY_BG.get(urgency, "FFFFFF"))


def urgency_cell_fill(urgency: str) -> PatternFill:
    return fill(URGENCY_CELL_BG.get(urgency, "FFFFFF"))


def status_cell_fill(status: str) -> PatternFill:
    return fill(STATUS_CELL_BG.get(status, "FFFFFF"))


# ── Fontes ────────────────────────────────────────────────────────────────────

def section_font() -> Font:
    return Font(bold=True, color=SECTION_FG, size=13)


def col_header_font() -> Font:
    return Font(bold=True, color=COL_HEADER_FG, size=10)


def data_font() -> Font:
    return Font(size=10)


def bold_data_font() -> Font:
    return Font(bold=True, size=10)


# ── Bordas ────────────────────────────────────────────────────────────────────

def thin_border() -> Border:
    thin = Side(border_style="thin", color="D0D0D0")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


# ── Alinhamentos ──────────────────────────────────────────────────────────────

def center_align() -> Alignment:
    return Alignment(horizontal="center", vertical="center")


def left_align() -> Alignment:
    return Alignment(horizontal="left", vertical="center", wrap_text=True)
