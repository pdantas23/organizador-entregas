"""
ExcelWriter — escreve dados brutos em uma worksheet.

Responsabilidades:
  - Inserir cabeçalhos de seção (MANHÃ / TARDE).
  - Inserir cabeçalhos de colunas.
  - Inserir linhas de dados com estilo delegado ao Styler.
  - Configurar larguras de colunas e freeze de panes.

NÃO decide ordenação nem agrupamento: recebe os dados já organizados.
"""
from __future__ import annotations

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from excel import styler
from models.delivery import Delivery
from utils.constants import COL_WIDTHS, OUTPUT_COLUMNS, PERIOD_ICONS

NUM_COLS = len(OUTPUT_COLUMNS)


class ExcelWriter:
    """Cursor de escrita que avança linha a linha na worksheet."""

    def __init__(self, ws: Worksheet) -> None:
        self._ws = ws
        self._row = 1  # linha atual

    # ── API pública ───────────────────────────────────────────────────────────

    def write_section_header(self, period: str) -> None:
        icon = PERIOD_ICONS.get(period, "")
        label = f"  {icon}  {period.upper()}  "
        cell = self._ws.cell(row=self._row, column=1, value=label)

        self._ws.merge_cells(
            start_row=self._row, start_column=1,
            end_row=self._row, end_column=NUM_COLS,
        )
        cell.fill = styler.section_fill()
        cell.font = styler.section_font()
        cell.alignment = styler.center_align()
        self._ws.row_dimensions[self._row].height = 28
        self._row += 1

    def write_column_headers(self) -> None:
        for col_i, name in enumerate(OUTPUT_COLUMNS, 1):
            cell = self._ws.cell(row=self._row, column=col_i, value=name)
            cell.fill = styler.col_header_fill()
            cell.font = styler.col_header_font()
            cell.alignment = styler.center_align()
            cell.border = styler.thin_border()
        self._ws.row_dimensions[self._row].height = 18
        self._row += 1

    def write_delivery_row(self, delivery: Delivery) -> None:
        row_data = delivery.to_excel_row()

        for col_i, (col_name, value) in enumerate(zip(OUTPUT_COLUMNS, row_data), 1):
            cell = self._ws.cell(row=self._row, column=col_i, value=value)
            cell.border = styler.thin_border()

            if col_name == "Urgência":
                cell.fill = styler.urgency_cell_fill(delivery.urgency)
                cell.font = styler.bold_data_font()
                cell.alignment = styler.center_align()

            elif col_name == "Status":
                cell.fill = styler.status_cell_fill(delivery.status)
                cell.font = styler.data_font()
                cell.alignment = styler.center_align()

            elif col_name == "Data":
                cell.fill = styler.row_fill(delivery.urgency)
                cell.font = styler.data_font()
                cell.number_format = "DD/MM/YYYY"
                cell.alignment = styler.center_align()

            elif col_name == "Nº Placas":
                cell.fill = styler.row_fill(delivery.urgency)
                cell.font = styler.data_font()
                cell.alignment = styler.center_align()

            else:
                cell.fill = styler.row_fill(delivery.urgency)
                cell.font = styler.data_font()
                cell.alignment = styler.left_align()

        self._ws.row_dimensions[self._row].height = 16
        self._row += 1

    def write_empty_period_notice(self) -> None:
        cell = self._ws.cell(
            row=self._row, column=1,
            value="(nenhuma entrega neste período)",
        )
        cell.font = styler.data_font()
        cell.alignment = styler.left_align()
        self._row += 1

    def write_blank_row(self) -> None:
        self._row += 1

    def finalize(self) -> None:
        """Aplica configurações finais: larguras e freeze de panes."""
        for i, width in enumerate(COL_WIDTHS, 1):
            self._ws.column_dimensions[get_column_letter(i)].width = width
        # Congela as duas primeiras linhas (cabeçalho de seção + colunas)
        self._ws.freeze_panes = "A3"
